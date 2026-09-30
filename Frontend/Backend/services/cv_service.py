import os
import json
from io import BytesIO

from pypdf import PdfReader
from docx import Document
from groq import Groq


# -------------------------
# PDF TEXT EXTRACTION
# -------------------------

def extract_pdf_text(file_bytes: bytes) -> str:

    reader = PdfReader(BytesIO(file_bytes))

    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text.strip()


# -------------------------
# DOCX TEXT EXTRACTION
# -------------------------

def extract_docx_text(file_bytes: bytes) -> str:

    document = Document(BytesIO(file_bytes))

    text = []

    # Paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text.strip())

    # Tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())

            if row_text:
                text.append(" | ".join(row_text))

    return "\n".join(text).strip()


# -------------------------
# GROQ CLIENT
# -------------------------

client = Groq(
    api_key=os.getenv("GROQ_API_KEY")
)


# -------------------------
# JSON CLEANER
# -------------------------

def clean_json_response(content: str) -> dict:

    content = content.strip()

    if content.startswith("```json"):
        content = content[7:]

    elif content.startswith("```"):
        content = content[3:]

    if content.endswith("```"):
        content = content[:-3]

    return json.loads(content.strip())


# -------------------------
# AI CV ANALYSIS
# -------------------------

def analyze_cv_with_ai(cv_text: str) -> dict:

    prompt = f"""
You are an AI CV analysis and job recommendation system.

Analyze the candidate CV carefully.

Return ONLY valid JSON.
Do not include markdown.
Do not include explanations outside JSON.

Use exactly this structure:

{{
    "candidate_name": "",
    "professional_summary": "",
    "skills": [],
    "years_of_experience": 0,
    "experience": [
        {{
            "job_title": "",
            "company": "",
            "duration": "",
            "responsibilities": []
        }}
    ],
    "education": [
        {{
            "degree": "",
            "institution": "",
            "duration": ""
        }}
    ],
    "recommended_roles": [],
    "strengths": [],
    "skill_gaps": []
}}

Rules:
- Extract facts only when explicitly present in the CV.
- Never invent a company name, job title, date, degree, skill, or responsibility.
- Do not create employment experience from the candidate's name, email, filename, project, or education.
- If a field is uncertain, return an empty string instead of guessing.
- Keep each education qualification separate.
- Maximum 15 skills.
- Maximum 5 recommended roles.
- Maximum 8 strengths.
- Maximum 6 skill gaps.
- Maximum 4 responsibilities per experience entry.
- Keep values concise.
- Return valid JSON only.

CV TEXT:

{cv_text}
"""

    response = client.chat.completions.create(
        model="openai/gpt-oss-20b",

        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],

        temperature=0.2,
        max_completion_tokens=2500,
        reasoning_effort="low",
        include_reasoning=False,

        response_format={
            "type": "json_object"
        }
    )

    content = response.choices[0].message.content

    if not content:
        raise ValueError("AI returned an empty response.")

    return clean_json_response(content)